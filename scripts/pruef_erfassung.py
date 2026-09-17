"""Strukturprüfung der Erfassungsunterlagen.

    python3 scripts/pruef_erfassung.py

LibreOffice läuft in dieser Umgebung nicht, Word-Dateien lassen sich also nicht
rendern. Geprüft wird deshalb der Aufbau: ein Blatt je Raum, die richtigen
Nummern, die richtige Zeilenzahl, die Vorbelegung der schon erfassten Artikel
und die Deckungsgleichheit von Word und Excel.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
import build_erfassung as B

fehler, proben = [], 0


def pruefe(name, ok, zusatz=""):
    global proben
    proben += 1
    print(("  ok       " if ok else "  FEHLER   ") + name + ("" if ok else f"   {zusatz}"))
    if not ok:
        fehler.append(name)


def main():
    raeume, je_raum = B.lade()
    soll = {f"{r['Raumcode']}-{n:02d}" for r in raeume
            for n in range(1, B.zeilenzahl(r, len(je_raum.get(r["Raumcode"], []))) + 1)}

    print("== Türschilder ==")
    d = Document(os.path.join(B.AUSGABE, "T1_Tuerschilder.docx"))
    pruefe("ein Schild je Raum", len(d.tables) == len(raeume), f"{len(d.tables)}/{len(raeume)}")
    umbrueche = len(d.element.body.findall(".//" + qn("w:br")))
    pruefe("fünf Schilder je Blatt", umbrueche == (len(raeume) - 1) // 5, str(umbrueche))
    text = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    pruefe("jeder Raumcode kommt vor", all(r["Raumcode"] in text for r in raeume))
    pruefe("jeder Raumname kommt vor", all(r["Raumname"] in text for r in raeume))

    print("\n== Erfassungsblätter ==")
    d2 = Document(os.path.join(B.AUSGABE, "T2_Erfassungsblaetter.docx"))
    pruefe("ein Blatt je Raum", len(d2.tables) == len(raeume), f"{len(d2.tables)}")
    schief = []
    kopfwiederholt = 0
    for t, r in zip(d2.tables, raeume):
        code = r["Raumcode"]
        bereits = je_raum.get(code, [])
        erwartet = B.zeilenzahl(r, len(bereits))
        if len(t.rows) - 1 != erwartet:
            schief.append(f"{code}: {len(t.rows)-1} statt {erwartet} Zeilen")
        if t.rows[1].cells[0].text != f"{code}-01":
            schief.append(f"{code}: erste Nummer {t.rows[1].cells[0].text!r}")
        trPr = t.rows[0]._tr.find(qn("w:trPr"))
        if trPr is not None and trPr.find(qn("w:tblHeader")) is not None:
            kopfwiederholt += 1
        for i, a in enumerate(bereits, start=1):
            if a["Bezeichnung"][:30] not in t.rows[i].cells[1].text:
                schief.append(f"{code}-{i:02d}: Bezeichnung fehlt")
        if len(bereits) < erwartet and t.rows[len(bereits) + 1].cells[1].text.strip():
            schief.append(f"{code}: erste freie Zeile ist nicht leer")
    pruefe("Zeilenzahl, Vorbelegung und erste freie Zeile", not schief, str(schief[:4]))
    pruefe("Kopfzeile wiederholt sich auf Folgeseiten", kopfwiederholt == len(raeume), str(kopfwiederholt))
    hoehen = {round(r.height.cm, 2) for t in d2.tables for r in t.rows[1:]}
    pruefe("Zeilenhöhe 0,85 cm", hoehen == {0.85}, str(hoehen))
    ist = {t.rows[i].cells[0].text for t in d2.tables for i in range(1, len(t.rows))}
    pruefe("alle Nummern vorhanden", ist == soll,
           f"fehlt {sorted(soll-ist)[:3]} · zuviel {sorted(ist-soll)[:3]}")

    print("\n== Erfassungsliste ==")
    wb = load_workbook(os.path.join(B.AUSGABE, "T3_Erfassungsliste.xlsx"))
    pruefe("zwei Blätter", wb.sheetnames == ["Erfassung", "Anleitung"], str(wb.sheetnames))
    ws = wb["Erfassung"]
    kopf = [c.value for c in ws[1]]
    pruefe("Kopfzeile", kopf == ["Nummer", "Gebäude", "Ebene", "Raum", "Artikel",
                                 "Stückzahl", "Einheit", "Bemerkung"], str(kopf))
    nrs = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
    pruefe("keine doppelten Nummern", len(set(nrs)) == len(nrs))
    pruefe("deckungsgleich mit den Erfassungsblättern", set(nrs) == soll)
    pruefe("Stückzahl-Prüfung eingebaut", len(ws.data_validations.dataValidation) == 1)
    vorbelegt = [i for i in range(2, ws.max_row + 1) if ws.cell(row=i, column=5).value]
    anzahl_aktiv = sum(len(v) for v in je_raum.values())
    pruefe(f"{anzahl_aktiv} vorbelegte Zeilen", len(vorbelegt) == anzahl_aktiv, str(len(vorbelegt)))
    pruefe("vorbelegte Zeilen tragen eine Menge",
           all(ws.cell(row=i, column=6).value for i in vorbelegt))
    pruefe("Raumangaben überall gefüllt",
           all(ws.cell(row=i, column=2).value and ws.cell(row=i, column=4).value
               for i in range(2, ws.max_row + 1)))

    mehrseitig = [r["Raumcode"] for r in raeume
                  if B.zeilenzahl(r, len(je_raum.get(r["Raumcode"], []))) > 26]
    print(f"\n  {len(raeume)} Räume · {len(soll)} Nummern · {anzahl_aktiv} davon vorbelegt")
    print(f"  Räume mit zwei Blättern: {', '.join(mehrseitig)}")
    print(f"\n== Ergebnis: {proben} Prüfungen, {len(fehler)} Fehler ==")
    return 1 if fehler else 0


if __name__ == "__main__":
    sys.exit(main())
