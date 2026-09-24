"""Prüft Datenbasis und erzeugte Dateien. Vor jeder Übergabe laufen lassen.

    python3 scripts/pruefen.py

Gibt am Ende die Zahl der Beanstandungen aus und beendet sich mit Code 1, wenn es welche gibt.
"""
import sys, os, re, csv
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import lade_artikel, foto, AUSGABE, BASIS, eur, USt_SATZ, PAKETRABATT
from openpyxl import load_workbook

fehler, warnung = [], []
def F(t): fehler.append(t);  print("  FEHLER  ", t)
def W(t): warnung.append(t); print("  Hinweis ", t)
def OK(t): print("  ok      ", t)

print("\n== 1. Datenbasis ==")
alle = lade_artikel(nur_aktive=False)
aktiv = lade_artikel()
nummern = [a["ArtNr"] for a in alle]
if len(nummern) != len(set(nummern)):
    doppelt = {n for n in nummern if nummern.count(n) > 1}
    F(f"doppelte Artikelnummern: {sorted(doppelt)}")
else:
    OK(f"{len(alle)} Positionen, davon {len(aktiv)} aktiv, keine doppelten Nummern")

ohne_foto = [a["ArtNr"] for a in aktiv if not a["Foto"]]
fehlt = [f'{a["ArtNr"]} ({a["Foto"]})' for a in aktiv if a["Foto"] and not foto(a["Foto"])]
if fehlt: F(f"Fotodatei fehlt: {fehlt}")
else:     OK("alle Fotoverweise zeigen auf vorhandene Dateien")
if ohne_foto: W(f"{len(ohne_foto)} Positionen ohne Foto: {ohne_foto[:8]}")

ohne_preis = [a["ArtNr"] for a in aktiv if not a["Preis_netto"]]
if ohne_preis: W(f"{len(ohne_preis)} Positionen ohne Preis: {ohne_preis[:8]}")
else:          OK("alle aktiven Positionen haben einen Preis")
ohne_mass = [a["ArtNr"] for a in aktiv if not a.get("Maße")]
if ohne_mass: W(f"{len(ohne_mass)} Positionen ohne Maßangabe")
# Seit die internen Notizen in der Bemerkung stehen, muss dort mitgesucht werden -
# sonst verschwindet eine offene Stueckzahl still aus der Prüfung.
nachzaehlen = [a["ArtNr"] for a in aktiv
               if "nachzählen" in ((a.get("Mengenhinweis") or "") + (a.get("Bemerkung") or ""))]
if nachzaehlen: W(f"{len(nachzaehlen)} Positionen mit offener Stückzahl")

# Eine Ware, die nicht mehr im Haus steht (Aktiv = nein), darf nicht zugleich für den
# Katalog freigegeben sein - sonst reserviert jemand etwas, das es nicht mehr gibt.
# "entfällt" ist etwas anderes: das sind die alten K-Nummern, die durch das Raumschema
# ersetzt wurden. Sie fallen über Aktiv ohnehin aus jeder Ausgabe heraus.
_widerspruch = [a["ArtNr"] for a in alle
                if (a.get("Aktiv") or "ja").lower() == "nein"
                and (a.get("Im_Katalog") or "ja").lower() == "ja"]
if _widerspruch:
    F(f"Aktiv = nein, aber Im_Katalog = ja: {_widerspruch}")
else:
    OK("keine abgemeldete Position ist für den Katalog freigegeben")

# Zusammengefasste Positionen: die Post-it-Nummern der weiteren Räume müssen eindeutig
# bleiben und dürfen nicht zugleich als eigene Zeile existieren, sonst wird dieselbe Ware
# zweimal angeboten.
_weitere = []
for a in alle:
    _weitere += [n.strip() for n in (a.get("Weitere_ArtNr") or "").split(",") if n.strip()]
_kollision = sorted(set(_weitere) & set(nummern))
_mehrfach = sorted({n for n in _weitere if _weitere.count(n) > 1})
if _kollision:
    F(f"Nummer steht in Weitere_ArtNr und ist zugleich eine eigene Zeile: {_kollision}")
elif _mehrfach:
    F(f"Nummer mehrfach in Weitere_ArtNr: {_mehrfach}")
elif _weitere:
    OK(f"{len(_weitere)} zusammengefasste Post-it-Nummern, alle eindeutig")

# Positionen, die noch auf etwas warten, tragen "OFFEN:" am Anfang der Bemerkung. Sie hier
# aufzulisten ist zuverlässiger als eine Merkliste: der Hinweis steht an der Ware selbst und
# taucht bei jedem Lauf wieder auf, bis jemand ihn entfernt.
_offen = [a for a in aktiv if (a.get("Bemerkung") or "").strip().startswith("OFFEN:")]
if _offen:
    W(f"{len(_offen)} Positionen mit offenem Punkt:")
    for a in _offen:
        _text = a["Bemerkung"].split("OFFEN:", 1)[1].strip()
        _kurz = _text if len(_text) <= 96 else _text[:93] + "..."
        print(f"             {a['ArtNr']}  {_kurz}")
else:
    OK("keine Position mit offenem Punkt")

print("\n== 1b. Raumstammdaten ==")
import csv as _csv
_rp = os.path.join(BASIS, "daten", "raeume.csv")
_raeume = {r["Raumcode"]: r for r in _csv.DictReader(open(_rp, encoding="utf-8"), delimiter=";")}
_codes = [r["Raumcode"] for r in _csv.DictReader(open(_rp, encoding="utf-8"), delimiter=";")]
if len(set(_codes)) != len(_codes):
    F("doppelte Raumcodes in raeume.csv")
else:
    OK(f"{len(_raeume)} Räume, Codes eindeutig")
_ohne = sorted({a["Raumcode"] for a in alle if a.get("Raumcode") and a["Raumcode"] not in _raeume})
if _ohne:
    F("Artikel verweisen auf unbekannte Raumcodes: " + ", ".join(_ohne))
else:
    OK("alle Artikel verweisen auf einen bekannten Raum")
_schief = [a["ArtNr"] for a in alle if a.get("Raumcode") in _raeume
           and a["Raum"] != _raeume[a["Raumcode"]]["Raumname"]]
if _schief:
    F(f"Raumname weicht vom Raumcode ab: {', '.join(_schief[:8])}")
else:
    OK("Raumname und Raumcode passen zusammen")
_falsch = [a["ArtNr"] for a in aktiv if not a["ArtNr"].startswith(a["Raumcode"] + "-")]
if _falsch:
    F(f"Artikelnummer passt nicht zum Raumcode: {', '.join(_falsch[:8])}")
else:
    OK("jede Artikelnummer trägt ihren Raumcode")

print("\n== 2. Fotoindex ==")
idx = os.path.join(BASIS, "daten", "fotos_index.csv")
if not os.path.exists(idx):
    F("daten/fotos_index.csv fehlt – ohne Index werden Fotonummern neu vergeben")
else:
    rows = list(csv.DictReader(open(idx, encoding="utf-8"), delimiter=";"))
    nums = [r["Foto"] for r in rows]
    if len(nums) != len(set(nums)): F("doppelte Fotonummern im Index")
    else: OK(f"{len(rows)} Fotos im Index, Nummern eindeutig")
    dateien = {x[:-4] for x in os.listdir(os.path.join(BASIS, "fotos")) if x.endswith(".jpg")}
    verwaist = sorted(dateien - set(nums))
    if verwaist: W(f"Fotodateien ohne Indexeintrag: {verwaist[:8]}")

print("\n== 3. Arbeitsmappe ==")
p = os.path.join(AUSGABE, "01_Artikelstamm_kikripp.xlsx")
if not os.path.exists(p):
    F("01_Artikelstamm_kikripp.xlsx fehlt – erst erzeugen")
else:
    wb = load_workbook(p)
    erwartet = ["Anleitung", "Artikelstamm", "Verkaufsübersicht", "Rechnungen (DATEV)",
                "Kasse", "Design"]
    if wb.sheetnames != erwartet: F(f"Blätter weichen ab: {wb.sheetnames}")
    else: OK("alle sechs Blätter vorhanden")
    st, vu = wb["Artikelstamm"], wb["Verkaufsübersicht"]
    kopf = {c.column_letter: c.value for c in st[2] if c.value}
    # Zeile 8 „davon teilverkauft“ ist dazugekommen, alles darunter rückt eins nach unten.
    ERW = {5:["ArtNr"], 6:["Status"], 7:["Status"], 8:["Status"], 9:["Status"],
           10:["Restmenge"], 11:["Positionswert_netto"], 13:["Verkauft_Menge"],
           14:["Verkauft_Menge","Verkaufspreis_netto"],
           17:["Zahlung","Verkauft_Menge","Verkaufspreis_netto"],
           21:["Status","Rechnungsnr"], 22:["Status","Zahlung"],
           23:["Status","Abholtermin"], 24:["Preis_netto"], 25:["Maße"]}
    schief = []
    for z, erw in ERW.items():
        ist = list(dict.fromkeys(re.findall(r"Artikelstamm!\$([A-Z]{1,2})\$",
                                            vu.cell(row=z, column=2).value or "")))
        if [kopf.get(x) for x in ist] != erw:
            schief.append((z, vu.cell(row=z, column=1).value, [kopf.get(x) for x in ist], erw))
    if schief:
        for z, label, ist, erw in schief:
            F(f"Verkaufsübersicht Z{z} „{label}“ greift auf {ist} statt {erw}")
    else:
        OK(f"alle {len(ERW)} Kennzahlen greifen auf die richtigen Spalten")

    ZAHLEN = ["Preis_netto", "Anschaffungswert_netto", "Verkaufspreis_netto",
              "Menge", "Verkauft_Menge"]
    sp = {v: k for k, v in kopf.items()}
    text_in_zahl = []
    for name in ZAHLEN:
        if name not in sp: continue
        from openpyxl.utils import column_index_from_string
        c = column_index_from_string(sp[name])
        for r in range(3, st.max_row + 1):
            v = st.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip():
                text_in_zahl.append(f"{name} Zeile {r}: {v!r}")
    if text_in_zahl: F(f"Text in Zahlenspalten (Formeln rechnen nicht): {text_in_zahl[:5]}")
    else: OK("alle Geld- und Mengenfelder sind Zahlen")

print("\n== 4. Ausgabedateien ==")
# Der PDF-Katalog ist entfallen; der Webkatalog auf kikripp.de hat ihn abgelöst.
for name, mindest in (("01_Artikelstamm_kikripp.xlsx", 20),
                      ("04_Webkatalog_MOCKUP.html", 500), ("06_Webkatalog_geschuetzt.html", 500),
                      ("katalog_import.json", 10),
                      ("A1_Webshop_einrichten.pdf", 20), ("A2_Artikel_verwalten.pdf", 20),
                      ("A3_Fotografieren.pdf", 20), ("A4_Arbeitsanweisung.pdf", 20),
                      ("T1_Tuerschilder.docx", 10), ("T2_Erfassungsblaetter.docx", 20),
                      ("T3_Erfassungsliste.xlsx", 20)):
    pf = os.path.join(AUSGABE, name)
    if not os.path.exists(pf): F(f"{name} fehlt")
    elif os.path.getsize(pf) // 1024 < mindest: F(f"{name} ist auffällig klein")
    else: OK(f"{name} ({os.path.getsize(pf)//1024} KB)")

print("\n== 5. Summen ==")
gesamt = sum(a["Positionswert"] for a in aktiv)
print(f"  Gesamtwert netto  {eur(gesamt)}")
print(f"  zzgl. {int(USt_SATZ*100)} % USt    {eur(gesamt*USt_SATZ)}")
print(f"  brutto            {eur(gesamt*(1+USt_SATZ))}")
print(f"  Paketpreis netto  {eur(gesamt*(1-PAKETRABATT))}  (−{int(PAKETRABATT*100)} %)")
for k in ("A", "B", "C"):
    t = [a for a in aktiv if a["Wertklasse"] == k]
    print(f"  Wertklasse {k}      {len(t):>3} Positionen · {eur(sum(x['Positionswert'] for x in t))}")

print(f"\n== Ergebnis: {len(fehler)} Fehler, {len(warnung)} Hinweise ==")
sys.exit(1 if fehler else 0)
